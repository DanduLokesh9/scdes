"""The Workflow Helper — Appendix H made kinetic.

Carries one Registry entry gate by gate: shows the gate's checklist, flashes the
appendices due at that gate, renders them as pre-filled forms, and on save
writes the values into a real `.xlsx`, advances the stage, and routes the gate
decision to the Council when the project's risk tier demands it.

The gate → appendix map is **derived** by reading Appendix H's own requirement
text, not hardcoded. Derived entries are marked so OT can confirm them in
Configure — which is the right shape anyway: the map is an operational parameter,
not a fact about the world.
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field as dc_field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app import config as config_mod
from app import registry as registry_mod
from app import scoring
from app.audit import CORPUS
from app.authz import Actor, Decision, Target, guard
from app.ingest_xlsx import APPENDIX_DIR, Block, parse_appendix, parse_all

GATE_SHEET = re.compile(r"gate\s*(\d)", re.I)

#: Instruments named in prose rather than by letter.
INSTRUMENT_NAMES = {
    "use case taxonomy": "A",
    "risk classification": "B",
    "risk matrix": "B",
    "system registry": "C",
    "ai system registry": "C",
    "data readiness": "D",
    "vendor ai disclosure": "E",
    "vendor disclosure": "E",
    "community impact": "F",
    "bias test": "F",
    "project intake": "G",
    "intake form": "G",
    "gate review checklist": "H",
    "capability survey": "I",
    "contract capability": "I",
    "evaluation methodology": "J",
    "stopping rule": "J",
    "model card": "K",
    "incident report": "L",
    "table of authorities": "M",
    "council operating procedures": "N",
}

_APPENDIX_REF = re.compile(r"appendix\s+([A-N])\b", re.I)


@dataclass
class GateRequirement:
    number: str
    text: str
    group: str = ""
    status_cell: str = ""
    notes_cell: str = ""
    appendices: list[str] = dc_field(default_factory=list)


@dataclass
class GateSpec:
    gate: int
    title: str
    sheet: str
    requirements: list[GateRequirement] = dc_field(default_factory=list)
    appendices: list[str] = dc_field(default_factory=list)
    derived: bool = True

    @property
    def name(self) -> str:
        return registry_mod.GATE_NAMES.get(self.gate, self.title)


def _appendices_in(text: str) -> list[str]:
    found = {m.group(1).upper() for m in _APPENDIX_REF.finditer(text)}
    lowered = text.lower()
    for phrase, letter in INSTRUMENT_NAMES.items():
        if phrase in lowered:
            found.add(letter)
    return sorted(found)


def derive_gate_map() -> dict[int, GateSpec]:
    """Read Appendix H and work out which instruments each gate pulls in."""
    schema = parse_appendix(APPENDIX_DIR / "Appendix_H_Gate_Review_Checklists.xlsx")
    gates: dict[int, GateSpec] = {}

    for sheet in schema.sheets:
        m = GATE_SHEET.search(sheet.name)
        if not m:
            continue
        gate = int(m.group(1))
        spec = GateSpec(gate=gate, title=sheet.name, sheet=sheet.name)

        for block in sheet.blocks:
            if block.kind != "checklist":
                continue
            status_col = notes_col = None
            for col in block.columns:
                if col.lower() == "status":
                    status_col = col
                elif col.lower().startswith("notes"):
                    notes_col = col

            for row in block.sample_rows:
                number = (row.get("#") or "").strip()
                text = (row.get("Requirement") or "").strip()
                if not text or text.lower() == "requirement":
                    continue
                refs = _appendices_in(text)
                spec.requirements.append(GateRequirement(
                    number=number, text=text, group=row.get("_group", ""),
                    appendices=refs,
                ))
                spec.appendices.extend(refs)

        # Field cells for Status / Notes come from the parsed field list, which
        # carries the real cell addresses for writeback.
        cells: dict[str, dict[str, str]] = {}
        for block in sheet.blocks:
            if block.kind != "checklist":
                continue
            for field in block.fields:
                key = field.label.split(" — ")[0]
                col = field.label.split(" — ")[-1].lower()
                slot = cells.setdefault(key, {})
                if col == "status":
                    slot["status"] = field.value_cell
                elif col.startswith("notes"):
                    slot["notes"] = field.value_cell
        for req in spec.requirements:
            key_variants = [req.number, f"{req.group} · {req.number}" if req.group else ""]
            for key in key_variants:
                if key and key in cells:
                    req.status_cell = cells[key].get("status", "")
                    req.notes_cell = cells[key].get("notes", "")
                    break

        spec.appendices = sorted(set(spec.appendices))
        gates[gate] = spec

    return dict(sorted(gates.items()))


_gate_map: dict[int, GateSpec] | None = None


def gate_map(rebuild: bool = False) -> dict[int, GateSpec]:
    global _gate_map
    if _gate_map is None or rebuild:
        _gate_map = derive_gate_map()
    return _gate_map


def gate_map_summary() -> list[dict[str, Any]]:
    return [
        {"gate": spec.gate, "name": spec.name,
         "requirements": len(spec.requirements),
         "appendices": spec.appendices, "derived": spec.derived}
        for spec in gate_map().values()
    ]


# ------------------------------------------------------------------- rendering

@dataclass
class RenderedStep:
    project_id: str
    project_name: str
    gate: int
    gate_name: str
    risk_band: str
    council_required: bool
    requirements: list[dict[str, Any]] = dc_field(default_factory=list)
    due_appendices: list[dict[str, Any]] = dc_field(default_factory=list)
    form: list[dict[str, Any]] = dc_field(default_factory=list)
    form_appendix: str = ""
    prefilled: dict[str, Any] = dc_field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {
            "project_id": self.project_id, "project_name": self.project_name,
            "gate": self.gate, "gate_name": self.gate_name,
            "risk_band": self.risk_band, "council_required": self.council_required,
            "requirements": self.requirements,
            "due_appendices": self.due_appendices,
            "form": self.form, "form_appendix": self.form_appendix,
            "prefilled": self.prefilled,
        }


#: Which appendix a gate's on-screen form should render first.
PRIMARY_FORM = {0: "G", 1: "D", 2: "J", 3: "K", 4: "J", 5: "J"}

_APPENDIX_FILES: dict[str, Path] | None = None


def appendix_files() -> dict[str, Path]:
    global _APPENDIX_FILES
    if _APPENDIX_FILES is None:
        _APPENDIX_FILES = {}
        for path in APPENDIX_DIR.glob("Appendix_*.xlsx"):
            letter = path.name.split("_")[1].upper()
            _APPENDIX_FILES[letter] = path
    return _APPENDIX_FILES


def _prefill_value(field, project: registry_mod.Project) -> Any:
    label = field.label.lower()
    risk = project.risk()
    if "project name" in label or "system name" in label:
        return project.name
    if "registry id" in label:
        return project.registry_id
    if "risk classification" in label:
        return risk.band
    if "bureau" in label or "program" in label and field.options:
        return project.bureau
    if "current gate" in label:
        return f"Gate {project.gate}"
    if "operational owner" in label:
        return project.operational_owner
    if "technical owner" in label:
        return project.technical_owner
    if "data owner" in label:
        return project.data_owner
    if "submission date" in label or "review date" in label or "assessment date" in label:
        return datetime.now(timezone.utc).date().isoformat()
    if "problem" in label or "description" in label:
        return project.description
    if "category" in label and field.options:
        return str(project.category)
    return None


def render_step(registry_id: str, gate: int | None = None) -> RenderedStep:
    project = registry_mod.load(registry_id)
    if project is None:
        raise KeyError(registry_id)
    gate = project.gate if gate is None else gate
    spec = gate_map().get(gate)
    risk = project.risk()

    step = RenderedStep(
        project_id=project.registry_id, project_name=project.name, gate=gate,
        gate_name=registry_mod.GATE_NAMES.get(gate, ""),
        risk_band=risk.band, council_required=risk.council_required,
    )

    if spec:
        step.requirements = [
            {"number": r.number, "text": r.text, "group": r.group,
             "appendices": r.appendices, "status_cell": r.status_cell,
             "notes_cell": r.notes_cell}
            for r in spec.requirements
        ]
        due = set(spec.appendices) | {
            a.split()[-1] for a in risk.required_appendices if a.startswith("Appendix")
        }
        step.due_appendices = [
            {"letter": letter,
             "title": _appendix_title(letter),
             "from_risk_band": letter in {
                 a.split()[-1] for a in risk.required_appendices},
             "from_checklist": letter in spec.appendices}
            for letter in sorted(due)
        ]

    letter = PRIMARY_FORM.get(gate, "G")
    step.form_appendix = letter
    path = appendix_files().get(letter)
    if path:
        schema = parse_appendix(path)
        fields: list[dict[str, Any]] = []
        for sheet in schema.sheets:
            if sheet.name.lower().startswith("instruction"):
                continue
            for block in sheet.blocks:
                if block.kind not in ("field_block",):
                    continue
                for field in block.fields:
                    value = _prefill_value(field, project)
                    fields.append({
                        "key": field.key, "label": field.label,
                        "control": field.control, "options": field.options,
                        "help": field.help_text, "sheet": field.sheet,
                        "cell": field.value_cell, "value": value,
                        "section": block.title,
                    })
                    if value is not None:
                        step.prefilled[field.key] = value
            if fields:
                break
        step.form = fields
    return step


def _appendix_title(letter: str) -> str:
    path = appendix_files().get(letter)
    if not path:
        return f"Appendix {letter}"
    name = path.stem.split("_", 2)[-1].replace("_", " ")
    return f"Appendix {letter} — {name}"


# ------------------------------------------------------------------- writeback

def instance_path(project: registry_mod.Project, letter: str) -> Path:
    template = appendix_files()[letter]
    return project.path / template.name


def save_step(registry_id: str, letter: str, values: dict[str, str],
              actor: Actor, *, advance_to: int | None = None,
              note: str = "") -> dict[str, Any]:
    """Write the form into the project's own copy of the appendix workbook.

    The adopted template is never modified — it is copied into the project
    folder on first write, so every project carries its own auditable instrument.
    """
    project = registry_mod.load(registry_id)
    if project is None:
        raise KeyError(registry_id)

    decision = guard(
        actor, Target.APPENDIX_INSTANCE, f"save_appendix_{letter}",
        owners=project.owners,
        detail={"registry_id": registry_id, "appendix": letter,
                "fields_written": len(values), "gate": project.gate},
    )
    if not decision.allowed:
        return {"decision": decision.reason, "written": False,
                "requires_council": decision.requires_council}

    template = appendix_files()[letter]
    target = instance_path(project, letter)
    if not target.exists():
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(template, target)

    schema = parse_appendix(template)
    by_key = {
        f.key: f
        for sheet in schema.sheets for block in sheet.blocks for f in block.fields
    }

    workbook = load_workbook(target)
    written: list[dict[str, str]] = []
    for key, value in values.items():
        field = by_key.get(key)
        if field is None or value in (None, ""):
            continue
        sheet = workbook[field.sheet] if field.sheet in workbook.sheetnames \
            else workbook.active
        sheet[field.value_cell] = value
        written.append({"key": key, "cell": f"{field.sheet}!{field.value_cell}",
                        "value": str(value)[:80]})
    workbook.save(target)
    workbook.close()

    project.appendix_instances[letter] = str(target.relative_to(CORPUS.parent))
    registry_mod.save(project)

    result: dict[str, Any] = {
        "written": True, "instance": str(target.relative_to(CORPUS.parent)),
        "cells": written, "count": len(written),
        "template_untouched": True,
    }

    if advance_to is not None:
        risk = project.risk()
        if risk.council_required:
            from app import council as council_mod
            proposal = council_mod.request_gate_decision(
                project=project, gate=advance_to, actor=actor,
                summary=note or f"Gate {advance_to} decision for {project.name}",
            )
            result["routed_to_council"] = proposal
            result["stage_advanced"] = False
            result["message"] = (
                f"{risk.band} risk — the Gate {advance_to} decision is the "
                f"Council's to make. Logged as {proposal['decision_id']} "
                f"awaiting a Tier {proposal['tier']} decision; the project stays "
                f"at Gate {project.gate}."
            )
        else:
            stage = {0: "concept", 1: "concept", 2: "pilot", 3: "production",
                     4: "production", 5: "production"}.get(advance_to, project.stage)
            _, project = registry_mod.advance_stage(
                registry_id, actor, stage=stage, gate=advance_to, note=note)
            result["stage_advanced"] = True
            result["message"] = (
                f"Low risk — bureau-level approval is sufficient, so the project "
                f"advanced to Gate {advance_to} ({registry_mod.GATE_NAMES.get(advance_to,'')})."
            )
    return result
