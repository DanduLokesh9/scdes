"""Read the exported framework the way the person adopting it would.

The client's note is a judgment about prose, not about structure::

    "Is the intent to additionally create a unified, proper document upon
     completion? I would think the download would have the 'good' version that
     could be adopted with the question audit trail at the end."

A unit test can assert that a heading exists. It cannot tell you whether the
paragraph under it reads like a policy or like a form somebody filled in. So
this fills in a realistic set of answers, builds the document, and prints it —
the framework part in full, so it can be read.

It also checks the properties that make it adoptable rather than a transcript:
the framework comes first, no question marks appear in it, the audit trail is
behind it, and nothing invented shows up in either.

Usage:  python -m tools.check_document           # print and check
        python -m tools.check_document --quiet   # check only
"""

from __future__ import annotations

import re
import sys
from datetime import date
from io import BytesIO

from docx import Document

from app import export, module_one, prose, versions
from app.authz import Actor, Role

WHO = Actor("sean.ot", "Dana Reed", Role.OT, title="Chief of Staff")

#: A twelve-person water district that has answered most of the module. Small
#: enough that the adaptive branches fire, complete enough that the document
#: has something to say in every section.
ANSWERS: dict[str, object] = {
    "org.kind": "district", "org.size": "u25", "org.adopter": "board",
    "org.functions": {"legal": "contracted", "it": "part", "security": "none",
                      "purchasing": "part", "records": "part",
                      "finance": "part", "hr": "none", "comms": "none"},
    "org.delegated": "yes", "org.open_records": "yes",
    "org.unusual": "We run a joint treatment plant with the next district "
                   "over, under an intergovernmental agreement.",

    "scope.covered": ["generative", "predictive", "recognition"],
    "scope.excluded": ["formulas", "spellcheck"],
    "scope.arbiter": "The District Manager",
    "scope.review": "standalone",

    "who.shape": "existing",
    "who.seats": [{"role": "District Manager"}, {"role": "Operations Lead"}],
    "who.tiebreak": "The District Manager",
    "who.consulted": ["legal", "it", "program"],
    "who.cadence": "quarterly",
    "who.without": ["free", "short_pilot"],
    "who.record": "The board's minutes, kept by the Clerk",
    "who.signs": "board", "who.amend": "same", "who.review": "annual",

    "risk.levels": "two",
    "risk.factors": {"consequence": "major", "public": "some",
                     "sensitive": "some", "irreversible": "major",
                     "delegated": "major", "disparate": "some",
                     "safety": "major", "news": "some"},
    "risk.tiers": {
        "routine": {"reviews": "The Operations Lead",
                    "written": "A line on the list of tools, and what it does",
                    "recheck": "annual", "nomeeting": "yes"},
        "elevated": {"reviews": "The District Manager, with counsel asked "
                                "first",
                     "written": "What it does, what information it touches, "
                                "who owns it, and the manual backup",
                     "recheck": "quarterly", "nomeeting": "no"}},
    "risk.worst": "yes",
    "risk.revisit": ["before_approval", "before_live", "annual",
                     "vendor_change", "incident"],

    "floor.final_action": "A service connection, a shut-off, a rate "
                          "determination, or a safety hold.",
    "floor.ai_finalises": "none",
    "floor.list_owner": "The Operations Lead",
    "floor.list_fields": ["name", "what", "owner", "data", "risk", "approved",
                          "offswitch"],
    "floor.list_checked": "biannual",
    "floor.unlisted": "violation",
    "floor.disclose_where": ["letters", "forms", "decisions"],
    "floor.reach_human": "Call the district office on the number printed on "
                         "every notice.",
    "floor.disclose_who": {"value": "The District Manager",
                           "approver": "The board"},
    "floor.disclose_home": "The shared drive, under Policies/AI",
    "floor.access_owner": "The Clerk",
    "floor.access_when": ["purchase", "live", "annual"],
    "floor.access_docs": "yes",
    "floor.fallback_scope": "all", "floor.fallback_tested": "annual",
    "floor.data_terms": ["no_training", "we_own", "returned", "no_sharing"],
    "floor.named": ["owner", "technical"],
    "floor.plain_scope": "all",
    "floor.plain_who": {"value": "The Operations Lead",
                        "approver": "The District Manager"},
    "floor.optional": ["reuse", "training"],

    "data.never": ["personal", "infrastructure", "personnel"],
    "data.where": "partly", "data.systems": "partly",
    "data.approver": "The District Manager",
    "data.retention": "silent", "data.prompts": "relied",

    "proc.today": "The board approves anything over $10,000; the District "
                  "Manager approves below that",
    "proc.cooperative": "regularly",
    "proc.terms": ["disclose", "notice", "accuracy", "return", "breach",
                   "accessibility"],
    "proc.checker": "Outside counsel",
    "proc.added_ai": "as_new", "proc.reuse": "yes", "proc.full_cost": "yes",
    "proc.cost_parts": ["licence", "setup", "training", "staff_time", "exit"],

    "watch.worth": "Fewer emergency call-outs, and the crew spends less time "
                   "on paperwork.",
    "watch.baseline": "sometimes",
    "watch.cadence": {"routine": {"every": "annual"},
                      "elevated": {"every": "quarterly"}},
    "watch.who": {"value": "The Operations Lead",
                  "where": "The tool register on the shared drive"},
    "watch.what": ["accuracy", "complaints", "security", "vendor"],
    "watch.failing": "pause",
    "watch.triggers": ["failed", "unsupported", "law", "contract",
                       "incident"],
    "watch.retire": {"value": "The District Manager",
                     "records": "Anything touching a service decision is kept "
                                "for seven years in the records system."},
    "watch.notice": "30",
    "watch.publish": ["annual_report"],

    "bad.existing": "yes", "bad.attach": "attach",
    "bad.report_to": "The Operations Lead, by phone or email",
    "bad.levels": {
        "minor": {"meaning": "A wrong result caught before it reached "
                             "anyone."},
        "serious": {"meaning": "It affected a customer or a decision."},
        "severe": {"meaning": "It affected safety, a rate, or the delegated "
                              "program."}},
    "bad.speed": {"minor": "3days", "serious": "sameday", "severe": "now"},
    "bad.stopper": "The Operations Lead or the District Manager",
    "bad.lookback": "yes", "bad.lookback_far": "needed",
    "bad.tell": ["board", "affected", "state", "federal"],
    "bad.writeup": {"value": "The Operations Lead",
                    "where": "The incident file on the shared drive"},

    "done.title": "Tidewater Water District AI Governance Framework",
    "done.signs": {"value": "board", "effective": "On signature"},
    "done.where": "both", "done.review": "annual",

    "words.03": "We deliberately kept the scope wide. It is easier to take "
                "something out later than to argue about whether it was ever "
                "in.",
}


def text_of(blob: bytes) -> tuple[str, str]:
    """The framework part and the appendix part, separately."""
    doc = Document(BytesIO(blob))
    parts = []
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    whole = "\n".join(p for p in parts if p.strip())
    marker = "Appendix — How this framework was decided"
    if marker in whole:
        front, back = whole.split(marker, 1)
        return front, back
    return whole, ""


def readable(blob: bytes) -> str:
    """The document in reading order, which the split above does not preserve
    because python-docx lists paragraphs and tables separately."""
    doc = Document(BytesIO(blob))
    body = doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    out = []
    for child in body.iterchildren():
        if child in paragraphs:
            said = paragraphs[child].text.strip()
            if said:
                out.append(said)
        elif child in tables:
            for row in tables[child].rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    out.append("    | " + " | ".join(cells))
    return "\n".join(out)


def main(argv: list[str]) -> int:
    quiet = "--quiet" in argv
    failures: list[str] = []

    def check(label: str, ok: bool, detail: str = "") -> None:
        print(f"  {'ok  ' if ok else 'FAIL'}  {label:<50} {detail}")
        if not ok:
            failures.append(label)

    for key, value in ANSWERS.items():
        versions.answer(key, value, WHO)

    blob = export.build("Tidewater Water District", today=date(2026, 9, 2))
    front, back = text_of(blob)

    if not quiet:
        print("=" * 74)
        print(readable(blob))
        print("=" * 74)
        print()

    print("the framework reads as a framework")
    check("it is a document Word can open", len(blob) > 10_000,
          f"{len(blob):,} bytes")
    check("the framework comes before the audit trail", bool(back),
          "appendix found at the back")

    # The heart of his note: the adoptable part must not read as a form. The
    # precise claim is that none of the module's questions appears in it,
    # tested against the questions themselves rather than against a pattern —
    # a pattern for "looks like a question" also matches "1.1  These rules
    # cover…", which is a clause and is meant to be numbered.
    prompts = [q.prompt for q in module_one.all_questions()]
    leaked = [p for p in prompts if p in front]
    check("no question from the module appears in the framework",
          not leaked, leaked[0][:46] if leaked else f"none of {len(prompts)}")
    # The gaps table quotes the unanswered ones on purpose — that is what a
    # gap is — so it is the one place a question mark is expected.
    before_gaps = front.split("What is not settled yet")[0]
    stray = [line.strip() for line in before_gaps.splitlines()
             if line.strip().endswith("?")
             and not line.strip().startswith(("How to read", "Why this"))]
    check("and no stray question marks before the gaps section",
          not stray, stray[0][:46] if stray else "clean")

    print("\nit says what they decided")
    for phrase in ("These rules cover",
                   "Responsibility for AI decisions sits with",
                   "This organization applies two levels of scrutiny",
                   "No AI tool finalizes an action",
                   "A single register of every AI tool"):
        check(f"“{phrase[:44]}…”", phrase in front)

    print("\nin their words, not ours")
    check("their own paragraph is carried verbatim",
          "easier to take something out later" in front)
    check("their title is used", "Tidewater Water District AI Governance"
          in prose.title_of(versions.state()["working"], "Tidewater"))

    print("\nand it adapts to who they are")
    check("no cybersecurity office is named", "cybersecurity" not in
          front.lower())
    check("the delegated program is carried",
          "delegated" in front.lower())
    check("a federal partner is in the telling list",
          "federal partner" in front.lower())

    print("\nthe audit trail is behind it, whole")
    for number in ("1.1", "5.3", "6.4c", "9.10", "10.8"):
        check(f"question {number} is in the appendix", number in back)
    check("with who answered", "Dana Reed" in back)

    print("\nnobody else is in it")
    for marker in ("SCDES", "South Carolina Department", "Appendix C",
                   "Operations Manual"):
        check(f"no mention of {marker}", marker not in front + back)

    print()
    if failures:
        print(f"FAIL — {len(failures)}: " + "; ".join(failures[:4]))
        return 1
    print("PASS — the framework reads as a framework, with the trail behind it")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))


