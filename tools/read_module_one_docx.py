"""Print the client's Module One document as plain text, step by step.

His wording is the specification — the prompts, the response sets and the
recommendations are all his, and the standing instruction is to stick to what
he wrote. This prints a named slice of the document so a build of the later
steps quotes the source rather than a summary of it.

Usage:  python -m tools.read_module_one_docx            # the whole thing
        python -m tools.read_module_one_docx "STEP 5"   # from that heading on
        python -m tools.read_module_one_docx "STEP 5" "STEP 7"
"""

from __future__ import annotations

import sys
from pathlib import Path

import docx

#: Where the client's document sits on this machine.
SOURCE = Path.home() / "Downloads" / "Governance_Framework_Builder_Module_One.docx"


def _blocks(path: Path):
    """Yield (style, text) for paragraphs and flattened tables, in order."""
    document = docx.Document(str(path))
    body = document.element.body
    paragraphs = {p._p: p for p in document.paragraphs}
    tables = {t._tbl: t for t in document.tables}
    for child in body.iterchildren():
        if child in paragraphs:
            para = paragraphs[child]
            text = para.text.strip()
            if text:
                yield para.style.name, text
        elif child in tables:
            for row in tables[child].rows:
                cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
                if any(cells):
                    yield "Table", " | ".join(cells)


def main(argv: list[str]) -> int:
    if not SOURCE.exists():
        print(f"not found: {SOURCE}")
        return 1
    start = argv[0].lower() if argv else None
    stop = argv[1].lower() if len(argv) > 1 else None

    printing = start is None
    for style, text in _blocks(SOURCE):
        low = text.lower()
        if not printing and start and low.startswith(start):
            printing = True
        elif printing and stop and low.startswith(stop):
            break
        if printing:
            tag = "T" if style == "Table" else style[:9]
            print(f"[{tag:<9}] {text}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
