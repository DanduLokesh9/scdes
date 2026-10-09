"""What shape is the real adopted framework?

The client, on the generated output: "the Governance Framework document that
is supposed to be generated as the valuable output needs to read like a formal
document which could be handed to someone and signed off on", and later, with
two photographs of an adopted framework attached: "The questions we are asking
should provide the inputs needed to create a draft that matches this."

So the target is not a document I design. It is the one IIA already writes.
This dumps the reference framework's own conventions — its front matter, its
heading levels and numbering, its styles, its section order, whether it has a
contents page, a definitions clause, a signature block, a revision table —
so the generator can be built to match rather than to approximate.

Usage:  python -m tools._read_structure [path] [--full]
"""

from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.shared import Pt

REFERENCE = (Path(__file__).resolve().parent.parent / "corpus" / "framework"
             / "SCDES_AI_Governance_Framework_Final.docx")


def describe(path: Path, full: bool = False) -> None:
    document = docx.Document(str(path))
    print(f"=== {path.name}")

    section = document.sections[0]
    print(f"\npage      : {section.page_width.inches:.2f} x "
          f"{section.page_height.inches:.2f} in")
    print(f"margins   : top {section.top_margin.inches:.2f}  "
          f"bottom {section.bottom_margin.inches:.2f}  "
          f"left {section.left_margin.inches:.2f}  "
          f"right {section.right_margin.inches:.2f}")

    for name, part in (("header", section.header), ("footer", section.footer)):
        lines = [p.text.strip() for p in part.paragraphs if p.text.strip()]
        print(f"{name:<10}: {lines or '(empty)'}")

    used: dict[str, int] = {}
    for para in document.paragraphs:
        used[para.style.name] = used.get(para.style.name, 0) + 1
    print("\nstyles in use:")
    for style, count in sorted(used.items(), key=lambda kv: -kv[1]):
        print(f"  {count:>4}  {style}")

    body = document.element.body
    print(f"\ntables    : {len(document.tables)}")
    for i, table in enumerate(document.tables[:6], 1):
        head = " | ".join(c.text.strip()[:26] for c in table.rows[0].cells)
        print(f"  t{i}  {len(table.rows)}x{len(table.columns)}  {head}")

    print("\n--- the spine (headings, and the first line under each)")
    last_was_heading = False
    shown = 0
    for para in document.paragraphs:
        text = " ".join(para.text.split())
        style = para.style.name
        if style.startswith("Heading") or style in ("Title", "Subtitle"):
            depth = 0 if style in ("Title", "Subtitle") else int(
                style.split()[-1]) if style.split()[-1].isdigit() else 1
            print(f"\n{'  ' * depth}[{style}] {text}")
            last_was_heading = True
            shown = 0
        elif text and (full or (last_was_heading and shown < 2)):
            print(f"{'      '}{text[:120]}")
            shown += 1
            if shown >= 2:
                last_was_heading = False

    # Numbering: is it in the text, or applied by a list style?
    numbered = [p.text.split()[0] for p in document.paragraphs
                if p.text.strip() and p.text.split()[0].rstrip(".").replace(
                    ".", "").isdigit()]
    print(f"\nparagraphs whose first token is a number: {len(numbered)}")
    print(f"  e.g. {numbered[:14]}")


def main(argv: list[str]) -> int:
    full = "--full" in argv
    args = [a for a in argv if not a.startswith("--")]
    path = Path(args[0]) if args else REFERENCE
    if not path.is_file():
        print(f"no such file: {path}")
        return 1
    describe(path, full)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
