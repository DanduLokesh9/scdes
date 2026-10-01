"""Print an appendix workbook's columns, so a module can be built from it
rather than from a guess about it.

The client on the vendor registry: "1 should be very easy to stand up bc you
already have a template as Appendix C." Worth reading the template before
agreeing.

Read-only.

Usage:  python -m tools._read_appendix C
        python -m tools._read_appendix Appendix_E_Vendor_AI_Disclosure.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

CORPUS = Path(__file__).resolve().parent.parent / "corpus" / "appendices"


def _find(which: str) -> Path | None:
    if (CORPUS / which).is_file():
        return CORPUS / which
    for path in sorted(CORPUS.glob("*.xlsx")):
        if path.name.lower().startswith(f"appendix_{which.lower()}_"):
            return path
    return None


def main(argv: list[str]) -> int:
    if not argv:
        for path in sorted(CORPUS.glob("*.xlsx")):
            print(f"  {path.name}")
        return 0

    path = _find(argv[0])
    if path is None:
        print(f"no appendix matching {argv[0]!r}")
        return 1

    print(f"{path.name}\n")
    book = load_workbook(path, data_only=True, read_only=True)
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(max_row=14, values_only=True))
        if not rows:
            continue
        print(f"--- sheet: {sheet.title}  ({sheet.max_row} rows, "
              f"{sheet.max_column} cols)")
        for i, row in enumerate(rows):
            cells = [str(c).strip() for c in row if c is not None
                     and str(c).strip()]
            if not cells:
                continue
            joined = " | ".join(cells)
            print(f"  r{i + 1}: {joined[:150]}")
        print()
    book.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
